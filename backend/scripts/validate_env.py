"""Phase 0 library validation.

Proves that every dependency in the required technology stack can be imported
and is functional at a basic level. Run with the project virtualenv:

    .venv/Scripts/python scripts/validate_env.py
"""
from __future__ import annotations

import importlib
import sys


def _report(name: str, ok: bool, detail: str = "") -> None:
    status = "OK " if ok else "FAIL"
    print(f"[{status}] {name}{(' — ' + detail) if detail else ''}")


def main() -> int:
    failures: list[str] = []

    checks = [
        ("fastapi", "fastapi"),
        ("uvicorn", "uvicorn"),
        ("pydantic", "pydantic"),
        ("pydantic_settings", "pydantic_settings"),
        ("sqlalchemy", "sqlalchemy"),
        ("pymysql", "pymysql"),
        ("alembic", "alembic"),
        ("fitz (PyMuPDF)", "fitz"),
        ("docx (python-docx)", "docx"),
        ("lxml", "lxml"),
        ("rapidfuzz", "rapidfuzz"),
        ("google.genai", "google.genai"),
        ("aiofiles", "aiofiles"),
        ("httpx", "httpx"),
    ]

    for label, module in checks:
        try:
            importlib.import_module(module)
            _report(label, True)
        except Exception as exc:  # noqa: BLE001
            _report(label, False, repr(exc))
            failures.append(label)

    # Functional smoke tests for the document-processing primitives.
    try:
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Liability cap is AED 100,000.")
        text = page.get_text().strip()
        assert "Liability cap" in text
        _report("PyMuPDF write+extract", True, f"extracted={text!r}")
    except Exception as exc:  # noqa: BLE001
        _report("PyMuPDF write+extract", False, repr(exc))
        failures.append("PyMuPDF functional")

    try:
        import docx

        d = docx.Document()
        p = d.add_paragraph("Supplier shall ")
        p.add_run("not be liable ")
        p.add_run("for indirect damages.")
        assert p.text == "Supplier shall not be liable for indirect damages."
        _report("python-docx multi-run paragraph", True, repr(p.text))
    except Exception as exc:  # noqa: BLE001
        _report("python-docx multi-run paragraph", False, repr(exc))
        failures.append("python-docx functional")

    try:
        from rapidfuzz import fuzz

        score = fuzz.token_set_ratio("within 30 days", "within  30   days")
        assert score == 100.0
        _report("rapidfuzz scoring", True, f"score={score}")
    except Exception as exc:  # noqa: BLE001
        _report("rapidfuzz scoring", False, repr(exc))
        failures.append("rapidfuzz functional")

    try:
        import lxml.etree  # noqa: F401
        from lxml import etree

        xml = '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:r><w:t>hi</w:t></w:r></w:p>'
        root = etree.fromstring(xml.encode())
        W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        assert root.find(f"{{{W}}}r/{{{W}}}t").text == "hi"
        _report("lxml OOXML namespace handling", True)
    except Exception as exc:  # noqa: BLE001
        _report("lxml OOXML namespace handling", False, repr(exc))
        failures.append("lxml functional")

    print(f"\nPython: {sys.version.split()[0]}")
    if failures:
        print(f"VALIDATION FAILED: {failures}")
        return 1
    print("VALIDATION PASSED: all required libraries import and function.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

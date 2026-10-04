"""Phase 8: real DOCX tracked changes (w:ins / w:del), incl. multi-run targets."""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document

from app.services.redlining.tracked_changes import (
    apply_tracked_edit,
    count_tracked_changes,
)

from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _build_docx(path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("The liability cap ")
    p.add_run("shall").bold = True
    p.add_run(" be mutual")
    p.add_run(" between the parties.")
    doc.add_paragraph("This paragraph stays untouched and must remain identical.")
    doc.save(str(path))


def test_tracked_change_applied_and_file_reopens(tmp_path: Path) -> None:
    src = tmp_path / "in.docx"
    _build_docx(src)

    doc = Document(str(src))
    edit = apply_tracked_edit(
        doc,
        "liability cap shall be mutual",
        "liability cap shall be mutual and reciprocal",
    )
    out = io.BytesIO()
    doc.save(out)
    data = out.getvalue()

    ins, dele = count_tracked_changes(data)
    assert (ins, dele) == (1, 1)
    assert edit.paragraph_index == 0

    # The file is a valid DOCX again.
    reopened = Document(io.BytesIO(data))
    assert len(reopened.paragraphs) >= 2

    # Inspect XML: delText carries the old words, ins carries the replacement.
    import zipfile
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    del_texts = "".join(t.text or "" for t in root.iter(f"{{{W}}}delText"))
    ins_texts = [
        "".join(t.text or "" for t in ins.iter(f"{{{W}}}t"))
        for ins in root.iter(f"{{{W}}}ins")
    ]
    assert del_texts == "liability cap shall be mutual"
    assert "liability cap shall be mutual and reciprocal" in ins_texts[0]


def test_bold_run_formatting_preserved_in_deletion(tmp_path: Path) -> None:
    src = tmp_path / "in.docx"
    _build_docx(src)
    doc = Document(str(src))
    apply_tracked_edit(doc, "liability cap shall be mutual", "mutual cap")
    out = io.BytesIO()
    doc.save(out)

    import zipfile
    with zipfile.ZipFile(io.BytesIO(out.getvalue())) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    # The bolded "shall" run must still be bold inside the deletion.
    bold_in_del = False
    for del_el in root.iter(f"{{{W}}}del"):
        for r in del_el.iter(f"{{{W}}}r"):
            if r.find(f"{{{W}}}delText") is not None and (
                r.find(f"{{{W}}}rPr/{{{W}}}b") is not None
            ):
                bold_in_del = True
    assert bold_in_del, "run formatting (bold) must survive the w:del wrap"


def test_untouched_paragraph_unchanged(tmp_path: Path) -> None:
    src = tmp_path / "in.docx"
    _build_docx(src)
    untouched_text = (
        "This paragraph stays untouched and must remain identical."
    )
    doc = Document(str(src))
    apply_tracked_edit(doc, "liability cap shall be mutual", "mutual cap")
    out = io.BytesIO()
    doc.save(out)
    reopened = Document(io.BytesIO(out.getvalue()))
    assert reopened.paragraphs[1].text == untouched_text


def test_target_must_occur_exactly_once(tmp_path: Path) -> None:
    src = tmp_path / "dup.docx"
    doc = Document()
    doc.add_paragraph("Fee is due. Fee is due. Fee is due again here now.")
    doc.save(str(src))
    doc = Document(str(src))
    with pytest.raises(ValueError):
        apply_tracked_edit(doc, "Fee is due.", "Amount is payable.")


def test_table_content_is_visible_and_editable(tmp_path: Path) -> None:
    """A value that lives only inside a table cell must be part of the document
    the redline verifies and edits - the viewer/extractor show table content, so
    the run tree the redline walks has to include it too (regression: the
    verification used ``document.paragraphs``, which silently omits tables)."""
    from app.services.redlining.tracked_changes import count_occurrences, paragraph_texts

    src = tmp_path / "tbl.docx"
    doc = Document()
    doc.add_paragraph("6. Operational Table")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Requirement"
    table.cell(1, 0).text = "Payment review period"
    table.cell(1, 1).text = "6 business days"
    doc.save(str(src))

    doc = Document(str(src))
    assert "6 business days" in paragraph_texts(doc)
    assert count_occurrences(doc, "6 business days") == 1

    edit = apply_tracked_edit(doc, "6 business days", "10 business days")
    out = io.BytesIO()
    doc.save(out)
    ins, dele = count_tracked_changes(out.getvalue())
    assert (ins, dele) == (1, 1)
    assert edit.target == "6 business days"


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(pytest.main([__file__, "-q"]))
